"""
Public document parsing (PDF / DOCX / XLSX / text / HTML) with pure-Python
parsers only.  No macros, no external viewers, bounded input.
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from datetime import datetime, timezone
from typing import Optional

from extractors.text_utils import strip_html
from extractors.url_utils import file_type_from_url

MAX_BYTES = 20 * 1024 * 1024
MAX_PDF_PAGES = 40


@dataclass
class DocumentText:
    text: str = ""
    title: str = ""
    author: str = ""
    organization: str = ""
    created: Optional[datetime] = None
    doc_type: str = ""
    pages: int = 0
    error: str = ""


def doc_type_from(content_type: str, url: str) -> str:
    ct = (content_type or "").lower()
    if "pdf" in ct:
        return "pdf"
    if "wordprocessingml" in ct:
        return "docx"
    if "msword" in ct:
        return "doc"
    if "spreadsheetml" in ct:
        return "xlsx"
    if "ms-excel" in ct:
        return "xls"
    if "presentationml" in ct:
        return "pptx"
    if "ms-powerpoint" in ct:
        return "ppt"
    if "html" in ct:
        return "html"
    if "csv" in ct:
        return "csv"
    if ct.startswith("text/"):
        return "txt"
    return file_type_from_url(url) or ""


def _pdf_date(raw: object) -> Optional[datetime]:
    s = str(raw or "")
    if s.startswith("D:"):
        s = s[2:]
    digits = "".join(ch for ch in s[:14] if ch.isdigit())
    if len(digits) >= 8:
        try:
            return datetime(int(digits[:4]), int(digits[4:6]), int(digits[6:8]), tzinfo=timezone.utc)
        except ValueError:
            return None
    return None


def _parse_pdf(content: bytes, out: DocumentText) -> None:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(content))
    out.pages = len(reader.pages)
    parts: list[str] = []
    for page in reader.pages[:MAX_PDF_PAGES]:
        try:
            parts.append(page.extract_text() or "")
        except Exception as exc:  # pragma: no cover - malformed page
            out.error = f"page extraction: {type(exc).__name__}"
    out.text = "\n".join(parts)
    meta = reader.metadata or {}
    try:
        out.title = str(meta.get("/Title") or "").strip()
        out.author = str(meta.get("/Author") or "").strip()
        out.organization = str(meta.get("/Creator") or meta.get("/Producer") or "").strip()
        out.created = _pdf_date(meta.get("/CreationDate"))
    except Exception:  # pragma: no cover
        pass


def _parse_docx(content: bytes, out: DocumentText) -> None:
    try:
        import docx  # type: ignore
    except ImportError:
        out.error = "python-docx not installed"
        return
    d = docx.Document(io.BytesIO(content))
    out.text = "\n".join(p.text for p in d.paragraphs)
    for table in d.tables[:50]:
        for row in table.rows:
            out.text += "\n" + " | ".join(c.text for c in row.cells)
    cp = d.core_properties
    out.title, out.author = cp.title or "", cp.author or ""
    if cp.created:
        out.created = cp.created.replace(tzinfo=timezone.utc) if cp.created.tzinfo is None else cp.created


def _parse_xlsx(content: bytes, out: DocumentText) -> None:
    try:
        import openpyxl  # type: ignore
    except ImportError:
        out.error = "openpyxl not installed"
        return
    wb = openpyxl.load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    cells: list[str] = []
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if v is not None:
                    cells.append(str(v))
                    if len(cells) >= 5000:
                        break
            if len(cells) >= 5000:
                break
    out.text = "\n".join(cells)
    props = wb.properties
    out.title, out.author = props.title or "", props.creator or ""


def extract_document(content: bytes, content_type: str, url: str) -> DocumentText:
    out = DocumentText(doc_type=doc_type_from(content_type, url))
    if not content:
        out.error = "empty content"
        return out
    if len(content) > MAX_BYTES:
        out.error = f"document too large ({len(content)} bytes)"
        return out
    try:
        if out.doc_type == "pdf" or content[:5] == b"%PDF-":
            out.doc_type = "pdf"
            _parse_pdf(content, out)
        elif out.doc_type == "docx":
            _parse_docx(content, out)
        elif out.doc_type == "xlsx":
            _parse_xlsx(content, out)
        elif out.doc_type in ("html", "htm"):
            page = strip_html(content.decode("utf-8", "ignore"))
            out.text, out.title = page.text, page.title
            out.author = page.meta.get("author", "")
        elif out.doc_type in ("txt", "csv", ""):
            out.text = content.decode("utf-8", "ignore")
        else:
            out.error = f"unsupported document type {out.doc_type}"
    except Exception as exc:
        out.error = f"{type(exc).__name__}: {exc}"[:200]
    out.text = out.text[:400_000]
    return out
